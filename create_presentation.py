import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette: Professional Clean White Academic Theme
    C_BG = RGBColor(255, 255, 255)          # #FFFFFF Pure White
    C_CARD = RGBColor(248, 250, 252)        # #F8FAFC Very Light Slate
    C_BORDER = RGBColor(226, 232, 240)      # #E2E8F0 Soft Border
    C_NAVY = RGBColor(30, 58, 138)          # #1E3A8A BVM Royal Navy Blue
    C_BLUE = RGBColor(37, 99, 235)          # #2563EB Vibrant Blue
    C_GOLD = RGBColor(217, 119, 6)          # #D97706 Warm Amber/Gold
    C_TEXT = RGBColor(15, 23, 42)           # #0F172A Dark Charcoal Text
    C_MUTED = RGBColor(71, 85, 105)         # #475569 Slate Muted Text
    C_LINE = RGBColor(203, 213, 225)        # #CBD5E1 Divider Line

    def add_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, slide_num, title, subtitle):
        # Top banner line
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.04))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_BLUE
        top_bar.line.fill.background()

        # Header text box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.733), Inches(1.2))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"SLIDE {slide_num}  •  MOVIELENS RECOMMENDER SYSTEM  •  BVM ENGINEERING COLLEGE"
        p0.font.name = "Arial"
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_BLUE

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = "Arial"
            p2.font.size = Pt(12)
            p2.font.color.rgb = C_MUTED

    def add_feature_card(slide, left, top, width, height, tag, tag_color, title, points):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = C_BORDER

        # Top accent strip on card
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = tag_color
        accent.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.18)
        tf.margin_bottom = Inches(0.12)

        p = tf.paragraphs[0]
        p.text = tag
        p.font.name = "Arial"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = tag_color

        p = tf.add_paragraph()
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_NAVY
        p.space_after = Pt(4)

        for pt in points:
            p = tf.add_paragraph()
            p.text = f"• {pt}"
            p.font.name = "Arial"
            p.font.size = Pt(10)
            p.font.color.rgb = C_TEXT
            p.space_after = Pt(3)

    # ==========================================
    # SLIDE 1: Title Slide (White Background)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_background(s1)

    # Top Navy Bar
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = C_NAVY
    top_bar.line.fill.background()

    # BVM Logo at Top Left (Crisp on White)
    bvm_logo = "Images/bvm_logo.png"
    if os.path.exists(bvm_logo):
        s1.shapes.add_picture(bvm_logo, Inches(0.9), Inches(0.5), Inches(1.4), Inches(1.5))

    # College & Department Header
    inst_box = s1.shapes.add_textbox(Inches(2.5), Inches(0.65), Inches(9.8), Inches(1.2))
    itf = inst_box.text_frame
    itf.word_wrap = True
    itf.margin_left = itf.margin_top = itf.margin_right = itf.margin_bottom = 0

    p = itf.paragraphs[0]
    p.text = "BIRLA VISHVAKARMA MAHAVIDYALAYA (ENGINEERING COLLEGE)"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p = itf.add_paragraph()
    p.text = "Department of Information Technology  •  An Autonomous Institution  •  Vallabh Vidyanagar"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = C_MUTED

    # Divider Line
    line1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.2), Inches(11.5), Inches(0.02))
    line1.fill.solid()
    line1.fill.fore_color.rgb = C_LINE
    line1.line.fill.background()

    # Title & Metadata box
    tbox = s1.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(7.0), Inches(4.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "FINAL PROJECT PRESENTATION"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    p = tf1.add_paragraph()
    p.text = "Movie Recommender System\non MovieLens Dataset"
    p.font.name = "Arial"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p = tf1.add_paragraph()
    p.text = "Knowledge-Based, Collaborative Filtering (KNN), SVD Matrix Factorization & Softmax DNN"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.color.rgb = C_MUTED

    # Presenter Card
    card_p = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(5.1), Inches(3.4), Inches(1.8))
    card_p.fill.solid()
    card_p.fill.fore_color.rgb = C_CARD
    card_p.line.color.rgb = C_BORDER

    tf_p = card_p.text_frame
    tf_p.word_wrap = True
    p = tf_p.paragraphs[0]
    p.text = "PRESENTED BY"
    p.font.name = "Arial"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf_p.add_paragraph()
    p.text = "Harsh Gondaliya"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p = tf_p.add_paragraph()
    p.text = "ID No: 23IT405"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    p = tf_p.add_paragraph()
    p.text = "B.Tech Information Technology"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED

    # Mentor Card
    card_m = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(5.1), Inches(3.4), Inches(1.8))
    card_m.fill.solid()
    card_m.fill.fore_color.rgb = C_CARD
    card_m.line.color.rgb = C_BORDER

    tf_m = card_m.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = "PROJECT MENTOR"
    p.font.name = "Arial"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf_m.add_paragraph()
    p.text = "Prof. Dhruvin R. Kotak"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p = tf_m.add_paragraph()
    p.text = "Assistant Professor"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    p = tf_m.add_paragraph()
    p.text = "Information Technology Department"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED

    # Title Image
    img_path1 = "Images/netflix-recommendation-s.jpeg"
    if os.path.exists(img_path1):
        s1.shapes.add_picture(img_path1, Inches(8.2), Inches(2.5), Inches(4.3), Inches(4.2))

    # ==========================================
    # SLIDE 2: Introduction to Recommender Systems
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_background(s2)
    add_header(s2, 2, "Introduction to Recommender Systems", "The algorithmic backbone of personalized digital media and content discovery")

    # 4 Feature Cards in a clean 2x2 Grid
    c_w = Inches(5.7)
    c_h = Inches(2.45)
    c_left1 = Inches(0.8)
    c_left2 = Inches(6.833)
    c_top1 = Inches(1.95)
    c_top2 = Inches(4.65)

    add_feature_card(
        s2, c_left1, c_top1, c_w, c_h,
        "DEFINITION & ESSENCE", C_BLUE,
        "What is a Recommender System?",
        [
            "Information filtering algorithms that predict user ratings, preferences, and affinities for unrated items.",
            "Transforms massive catalogs into personalized candidate sets tailored to individual taste profiles.",
            "Powers modern digital consumption: drives >80% of hours watched on Netflix and >70% of views on YouTube."
        ]
    )

    add_feature_card(
        s2, c_left2, c_top1, c_w, c_h,
        "MOTIVATION & RELEVANCE", C_GOLD,
        "Why Recommenders Matter",
        [
            "Solves 'Choice Paralysis': Users face thousands of options and quickly abandon platforms without personalization.",
            "Surfaces the Long-Tail: Uncovers hidden gems beyond mainstream blockbusters for diverse user tastes.",
            "Drives Business Value: Maximizes user retention, platform engagement, session duration, and satisfaction."
        ]
    )

    add_feature_card(
        s2, c_left1, c_top2, c_w, c_h,
        "MULTI-PARADIGM APPROACH", C_NAVY,
        "Core Methodologies Implemented",
        [
            "Knowledge-Based: Cold-start filtering via genre, watch count, and weighted ratings.",
            "Collaborative KNN: User-based & item-based similarity using Cosine distance.",
            "Matrix Factorization (SVD): 50 latent taste dimensions resolving extreme matrix sparsity.",
            "Softmax Deep Neural Network: Learned embedding layers with multi-class probabilities."
        ]
    )

    add_feature_card(
        s2, c_left2, c_top2, c_w, c_h,
        "PROJECT DELIVERABLES", C_BLUE,
        "Project Scope & Objectives",
        [
            "Benchmark Evaluation: Trained and evaluated on MovieLens 100K (100,000 ratings, 943 users, 1,682 movies).",
            "Comparative Insights: Evaluates cold-start handling, neighbor proximity, latent representations, and neural scoring.",
            "Full-Stack Interactive Dashboard: Real-time web application (Flask + JS/CSS) and CLI for live testing."
        ]
    )

    # ==========================================
    # SLIDE 3: Problem Statement & The Long-Tail
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_background(s3)
    add_header(s3, 3, "Problem Statement & The Long-Tail Challenge", "Why simple recommendation methods fail in real-world platforms")

    cbox3 = s3.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(6.2), Inches(4.8))
    tf3 = cbox3.text_frame
    tf3.word_wrap = True

    bullets3 = [
        ("Information Overload:", "Thousands of movies exist, making it difficult for users to find titles matching their specific taste."),
        ("The Cold-Start Problem:", "New users have zero watch history, preventing traditional collaborative models from making recommendations."),
        ("Data Sparsity:", "Over 93% of the movie-user interaction matrix is empty because users only rate a tiny fraction of catalog movies."),
        ("Popularity Bias (The Long Tail):", "Naive systems only push blockbusters. High-margin and niche movies in the 'long tail' remain unrecommended without matrix factorization.")
    ]
    for i, (head, desc) in enumerate(bullets3):
        p = tf3.add_paragraph() if i > 0 else tf3.paragraphs[0]
        p.text = f"• {head} {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.color.rgb = C_TEXT
        p.space_after = Pt(12)

    img_path3 = "Images/long-tail.png"
    if os.path.exists(img_path3):
        s3.shapes.add_picture(img_path3, Inches(7.4), Inches(2.2), Inches(5.1), Inches(3.8))

    # ==========================================
    # SLIDE 4: About the MovieLens 100K Dataset
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_background(s4)
    add_header(s4, 4, "MovieLens 100K Dataset & Data Insights", "Benchmark rating dataset collected by the GroupLens Research Project")

    cbox4 = s4.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(5.8), Inches(4.8))
    tf4 = cbox4.text_frame
    tf4.word_wrap = True

    bullets4 = [
        ("Dataset Scale:", "100,000 ratings from 943 users across 1,682 movies."),
        ("Rich Metadata:", "Includes 19 genre indicators, release dates, titles, and user demographics (age, gender, occupation)."),
        ("Active Participation:", "Each user in the dataset has rated at least 20 movies."),
        ("Rating Distribution:", "Users predominantly reward movies with 4-star ratings, followed by 3-star and 5-star ratings. Rarely do users rate below 2 stars.")
    ]
    for i, (head, desc) in enumerate(bullets4):
        p = tf4.add_paragraph() if i > 0 else tf4.paragraphs[0]
        p.text = f"• {head} {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.color.rgb = C_TEXT
        p.space_after = Pt(12)

    img_path4 = "Images/bar-plot.PNG"
    if os.path.exists(img_path4):
        s4.shapes.add_picture(img_path4, Inches(6.9), Inches(2.0), Inches(5.6), Inches(4.3))

    # ==========================================
    # SLIDE 5: Collaborative vs Content-Based Filtering
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_background(s5)
    add_header(s5, 5, "Recommendation Approaches: Content vs Collaborative", "Comparing how different algorithms make predictions")

    cbox5 = s5.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(5.8), Inches(4.8))
    tf5 = cbox5.text_frame
    tf5.word_wrap = True

    bullets5 = [
        ("Content-Based Filtering:", "Recommends items based on discrete characteristics (e.g. genre, actors, director). Suggests movies similar to what the user liked in the past."),
        ("Collaborative Filtering:", "Assumes users who agreed in the past will agree in the future. Uses user-item interaction profiles without needing item metadata."),
        ("Item-to-Item Similarity:", "Finds companion movies that have similar rating patterns across all viewers."),
        ("User-to-User Neighborhood:", "Finds like-minded peers and surfaces movies highly rated by those peers.")
    ]
    for i, (head, desc) in enumerate(bullets5):
        p = tf5.add_paragraph() if i > 0 else tf5.paragraphs[0]
        p.text = f"• {head} {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.color.rgb = C_TEXT
        p.space_after = Pt(12)

    img_path5 = "Images/coll-cont-pic.png"
    if os.path.exists(img_path5):
        s5.shapes.add_picture(img_path5, Inches(6.9), Inches(2.2), Inches(5.6), Inches(3.8))

    # ==========================================
    # SLIDE 6: Model 1 & 2: Knowledge-Based & KNN
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_background(s6)
    add_header(s6, 6, "Model 1 & 2: Knowledge-Based & K-Nearest Neighbors", "Instant cold-start resolution and memory-based collaborative filtering")

    cbox6 = s6.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(5.8), Inches(4.8))
    tf6 = cbox6.text_frame
    tf6.word_wrap = True

    bullets6 = [
        ("Knowledge-Based Recommender:", "Filters by genre, popularity (watch count > 300), and top average rating (> 4.0). Completely eliminates the Cold-Start hurdle for new visitors."),
        ("KNN Item-Based Filtering:", "Constructs a sparse [movies x users] matrix and computes Cosine Distances to find nearest movie neighbors in seconds."),
        ("KNN User-Based Filtering:", "Finds K closest users to target user ID, excludes already watched movies, and assigns distance-weighted scores to candidate films."),
        ("Dynamic Suggestor:", "Typo-tolerant string search suggests correct movie titles when users misspell query names.")
    ]
    for i, (head, desc) in enumerate(bullets6):
        p = tf6.add_paragraph() if i > 0 else tf6.paragraphs[0]
        p.text = f"• {head} {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.color.rgb = C_TEXT
        p.space_after = Pt(12)

    img_path6 = "Images/genre-bar-plot.PNG"
    if os.path.exists(img_path6):
        s6.shapes.add_picture(img_path6, Inches(6.9), Inches(2.1), Inches(5.6), Inches(4.2))

    # ==========================================
    # SLIDE 7: Model 3 & 4: SVD & Deep Neural Network
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_background(s7)
    add_header(s7, 7, "Model 3 & 4: SVD Matrix Factorization & Deep Learning", "Uncovering latent factors and capturing complex non-linear interactions")

    cbox7_left = s7.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8))
    tf7_l = cbox7_left.text_frame
    tf7_l.word_wrap = True

    p = tf7_l.paragraphs[0]
    p.text = "Singular Value Decomposition (SVD)"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    bullets7_l = [
        "Decomposes Utility Matrix A into U Σ V^T.",
        "Reduces dimensions to 50 latent factors representing hidden movie and user tastes.",
        "Approximates missing entries to predict unobserved ratings across the entire catalog.",
        "Overcomes data sparsity and handles the long-tail problem effectively."
    ]
    for b in bullets7_l:
        p = tf7_l.add_paragraph()
        p.text = f"• {b}"
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT
        p.space_after = Pt(8)

    cbox7_right = s7.shapes.add_textbox(Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8))
    tf7_r = cbox7_right.text_frame
    tf7_r.word_wrap = True

    p = tf7_r.paragraphs[0]
    p.text = "Softmax Deep Neural Network (DNN)"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    bullets7_r = [
        "Extracts latent features of users & movies using learned Embedding Layers.",
        "Passes representations through Dense layers with Dropout to prevent overfitting.",
        "Softmax output layer predicts rating probabilities (1 to 5 stars) for unseen items.",
        "Captures non-linear user-movie relationships that linear matrix factorization misses."
    ]
    for b in bullets7_r:
        p = tf7_r.add_paragraph()
        p.text = f"• {b}"
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 8: Live Web Application & Conclusion
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_background(s8)
    add_header(s8, 8, "Interactive Web Application & Key Takeaways", "Real-time evaluation dashboard and summary of achievements")

    cbox8 = s8.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(6.0), Inches(4.8))
    tf8 = cbox8.text_frame
    tf8.word_wrap = True

    bullets8 = [
        ("Unified Web Dashboard:", "Built an interactive modern UI (http://127.0.0.1:5050) allowing users to test all 4 models in real-time."),
        ("Multi-Paradigm Solution:", "Knowledge-based handles cold-start, KNN handles immediate similarity, SVD solves sparsity, and DNN models complex interactions."),
        ("Dynamic Movie Suggestor:", "Real-time fuzzy search prevents typos and delivers smooth user experience."),
        ("Conclusion:", "Combining statistical heuristics, neighborhood collaborative filtering, and latent space factorization delivers an optimal, highly scalable movie recommendation engine.")
    ]
    for i, (head, desc) in enumerate(bullets8):
        p = tf8.add_paragraph() if i > 0 else tf8.paragraphs[0]
        p.text = f"• {head} {desc}"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.color.rgb = C_TEXT
        p.space_after = Pt(12)

    img_path8 = "Images/pie.PNG"
    if os.path.exists(img_path8):
        s8.shapes.add_picture(img_path8, Inches(7.2), Inches(2.1), Inches(5.2), Inches(4.2))

    # ==========================================
    # SLIDE 9: Dedicated Thank You Slide (White Theme)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_background(s9)

    # Top Accent Bar
    top_bar9 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
    top_bar9.fill.solid()
    top_bar9.fill.fore_color.rgb = C_NAVY
    top_bar9.line.fill.background()

    # Center BVM Logo at top (High contrast on white!)
    if os.path.exists(bvm_logo):
        s9.shapes.add_picture(bvm_logo, Inches(5.7), Inches(0.6), Inches(1.8), Inches(1.9))

    # Thank you box
    ty_box = s9.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.5))
    ty_tf = ty_box.text_frame
    ty_tf.word_wrap = True

    p = ty_tf.paragraphs[0]
    p.text = "Thank You!"
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p.alignment = PP_ALIGN.CENTER

    p = ty_tf.add_paragraph()
    p.text = "Movie Recommender System on MovieLens Dataset"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.alignment = PP_ALIGN.CENTER

    p = ty_tf.add_paragraph()
    p.text = "Birla Vishvakarma Mahavidyalaya (Engineering College)  •  An Autonomous Institution"
    p.font.name = "Arial"
    p.font.size = Pt(12)
    p.font.color.rgb = C_MUTED
    p.alignment = PP_ALIGN.CENTER

    # Left Card: Student Details
    card_l = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.8), Inches(4.4), Inches(4.6), Inches(2.2))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = C_CARD
    card_l.line.color.rgb = C_BORDER

    tf_card_l = card_l.text_frame
    tf_card_l.word_wrap = True
    p = tf_card_l.paragraphs[0]
    p.text = "PRESENTED BY"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf_card_l.add_paragraph()
    p.text = "Harsh Gondaliya"
    p.font.name = "Arial"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p = tf_card_l.add_paragraph()
    p.text = "ID No: 23IT405"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    p = tf_card_l.add_paragraph()
    p.text = "Department of Information Technology\nBVM Engineering College"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED

    # Right Card: Mentor Details
    card_r = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(4.4), Inches(4.6), Inches(2.2))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = C_CARD
    card_r.line.color.rgb = C_BORDER

    tf_card_r = card_r.text_frame
    tf_card_r.word_wrap = True
    p = tf_card_r.paragraphs[0]
    p.text = "UNDER THE MENTORSHIP OF"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf_card_r.add_paragraph()
    p.text = "Prof. Dhruvin R. Kotak"
    p.font.name = "Arial"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p = tf_card_r.add_paragraph()
    p.text = "Assistant Professor"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    p = tf_card_r.add_paragraph()
    p.text = "Department of Information Technology\nBVM Engineering College"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED

    # Save presentation
    output_filename = "MovieLens_Recommender_System.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully as '{output_filename}' (White Academic Theme)!")

if __name__ == '__main__':
    create_deck()
